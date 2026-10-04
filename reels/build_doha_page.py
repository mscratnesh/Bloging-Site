"""Build dist/doha.html from निवेश_और_जीवन_के_100_दोहे_4.docx.

The page sits only in dist/ and no other page links to it yet.
Needs: pip install python-docx. Run again whenever the book changes.
"""
import html, pathlib, re
import docx

HERE = pathlib.Path(__file__).resolve().parent
BOOK = HERE / "निवेश_और_जीवन_के_100_दोहे_4.docx"
OUT = HERE.parent / "dist" / "doha.html"

d = docx.Document(str(BOOK))
dohas = []
for t in d.tables:
    head, mool, arth, bhav, niv, vya = (r.cells[1].text.strip() for r in t.rows)
    num, poet, lesson = (s.strip() for s in head.split("•"))
    dohas.append(dict(n=int(re.search(r"\d+", num).group()), poet=poet, lesson=lesson,
                      mool=mool, arth=arth, bhav=bhav, niv=niv, vya=vya))

paras = [p.text.strip() for p in d.paragraphs]
intro = paras[paras.index("भूमिका") + 1]
outro = paras[paras.index("निष्कर्ष") + 1]

e = html.escape


def couplet(text):
    # one line per half (split after ।), keeping the danda
    halves = [h.strip() for h in re.split(r"(?<=।)\s*", text) if h.strip()]
    return "<br>".join(e(h) for h in halves)


cards = []
for x in dohas:
    cards.append(f"""      <article class="doha" id="doha-{x['n']}" data-poet="{e(x['poet'])}">
        <header class="doha-head"><span class="doha-num">दोहा {x['n']}</span><span class="doha-poet">{e(x['poet'])}</span><h2>{e(x['lesson'])}</h2></header>
        <p class="couplet">{couplet(x['mool'])}</p>
        <details>
          <summary>अर्थ, भावार्थ, निवेश दोहा और व्याख्या</summary>
          <dl>
            <dt>अर्थ</dt><dd>{e(x['arth'])}</dd>
            <dt>भावार्थ</dt><dd>{e(x['bhav'])}</dd>
            <dt>निवेश दोहा</dt><dd class="couplet invest">{couplet(x['niv'])}</dd>
            <dt>व्याख्या</dt><dd>{e(x['vya'])}</dd>
          </dl>
        </details>
      </article>""")

poets = sorted({x["poet"] for x in dohas}, key=lambda p: -sum(y["poet"] == p for y in dohas))
chips = "".join(f'<button type="button" data-poet="{e(p)}">{e(p)} <small>{sum(y["poet"] == p for y in dohas)}</small></button>' for p in poets)

TOP = (HERE.parent / "calculators.html").read_text(encoding="utf-8")
strip = re.search(r'  <div class="top-strip">.*?</header>', TOP, re.S).group(0)
tail = re.search(r'  <aside class="site-disclaimer">.*?</footer>', TOP, re.S).group(0)

page = f"""<!doctype html>
<html lang="hi">
<head>
<meta charset="UTF-8"><meta name="viewport" content="width=device-width, initial-scale=1.0">
<meta name="description" content="कबीर, रहीम और तुलसीदास के 100 दोहे — हर दोहे का अर्थ, भावार्थ, निवेश दोहा और आधुनिक निवेश की व्याख्या।">
<link rel="canonical" href="https://letmoneyearn.in/doha.html">
<meta name="robots" content="index, follow">
<title>निवेश और जीवन के 100 दोहे | Let Money Earn</title>
<link rel="icon" type="image/png" href="https://raw.githubusercontent.com/mscratnesh/htmlSite/main/images/Let_Money_Earn_Logo_Cropped.png">
<meta property="og:type" content="website">
<meta property="og:site_name" content="Let Money Earn">
<meta property="og:title" content="निवेश और जीवन के 100 दोहे | Let Money Earn">
<meta property="og:description" content="कबीर, रहीम और तुलसीदास के दोहों से SIP, Compounding, Diversification और धैर्य की सीख।">
<meta property="og:url" content="https://letmoneyearn.in/doha.html">
<meta property="og:image" content="https://raw.githubusercontent.com/mscratnesh/htmlSite/main/images/Let_Money_Earn_Logo_Cropped.png">
<meta name="twitter:card" content="summary_large_image">
<meta name="twitter:title" content="निवेश और जीवन के 100 दोहे | Let Money Earn">
<meta name="twitter:description" content="कबीर, रहीम और तुलसीदास के दोहों से निवेश की सीख।">
<link rel="preconnect" href="https://fonts.googleapis.com"><link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
<link href="https://fonts.googleapis.com/css2?family=DM+Mono:wght@400;500&family=Fraunces:opsz,wght@9..144,500;9..144,600;9..144,700&family=Manrope:wght@400;500;600;700&family=Tiro+Devanagari+Hindi&family=Mukta:wght@400;500;600&display=swap" rel="stylesheet">
<link rel="stylesheet" href="styles.css">
<style>
  :root{{--deva:'Tiro Devanagari Hindi','Nirmala UI',serif;--hindi:'Mukta','Nirmala UI',var(--sans)}}
  .doha-hero{{padding:40px max(24px,calc((100% - 1180px)/2)) 28px;border-bottom:1px solid var(--line)}}
  .doha-hero h1{{max-width:820px;font:400 clamp(36px,5vw,58px)/1.15 var(--deva);margin:18px 0 0}}
  .doha-hero p{{max-width:640px;color:var(--muted);font:16px/1.7 var(--hindi);margin:22px 0 0}}
  .doha-tools{{padding:22px max(24px,calc((100% - 1180px)/2));display:flex;flex-wrap:wrap;gap:14px 24px;align-items:center;border-bottom:1px solid var(--line)}}
  .doha-tools input{{flex:1 1 260px;max-width:380px;border:0;border-bottom:1px solid var(--ink);background:transparent;padding:9px 0;outline:0;font:15px var(--hindi);color:var(--ink)}}
  .poet-chips{{display:flex;flex-wrap:wrap;gap:8px}}
  .poet-chips button,.open-all{{border:1px solid var(--line);background:transparent;padding:8px 15px;color:var(--muted);font:14px var(--hindi);cursor:pointer}}
  .poet-chips button small{{font:11px var(--mono)}}
  .poet-chips button.active,.poet-chips button:hover,.open-all:hover{{background:var(--ink);border-color:var(--ink);color:var(--paper)}}
  .open-all{{margin-left:auto}}
  .doha-count{{font:11px var(--mono);text-transform:uppercase;color:var(--muted)}}
  .doha-list{{padding:0 max(24px,calc((100% - 1180px)/2)) 80px;display:grid;grid-template-columns:repeat(2,minmax(0,1fr));gap:0 48px}}
  .doha{{border-top:1px solid var(--ink);padding:18px 0 30px}}
  .doha-head{{display:flex;flex-wrap:wrap;align-items:baseline;gap:4px 12px}}
  .doha-num{{font:500 11px var(--mono);letter-spacing:.06em;color:var(--coral-text)}}
  .doha-poet{{font:500 13px var(--hindi);color:var(--muted)}}
  .doha-head h2{{flex-basis:100%;font:600 21px/1.35 var(--hindi);margin:6px 0 0}}
  .couplet{{font:20px/1.75 var(--deva);margin:16px 0 0;color:var(--ink)}}
  details{{margin-top:14px}}
  summary{{cursor:pointer;font:13px var(--hindi);color:var(--coral-text);list-style:none}}
  summary::-webkit-details-marker{{display:none}}
  summary::before{{content:'+ ';font-family:var(--mono)}}
  details[open] summary::before{{content:'– '}}
  dl{{margin:12px 0 0}}
  dt{{font:500 10px var(--mono);letter-spacing:.08em;text-transform:uppercase;color:var(--muted);margin-top:14px}}
  dt:first-child{{margin-top:0}}
  dd{{margin:4px 0 0;font:15px/1.7 var(--hindi);color:var(--ink)}}
  dd.invest{{font:18px/1.75 var(--deva);color:var(--success);margin-top:4px}}
  .doha-note{{padding:0 max(24px,calc((100% - 1180px)/2)) 50px;max-width:820px;font:15px/1.7 var(--hindi);color:var(--muted)}}
  .doha-note h2{{font:600 22px var(--hindi);color:var(--ink);margin:0 0 8px}}
  .doha-disclaimer{{padding:22px max(24px,calc((100% - 1180px)/2));border-top:1px solid var(--line);font:13px/1.7 var(--hindi);color:var(--muted)}}
  .doha-disclaimer strong{{display:block;color:var(--ink);font:500 10px var(--mono);letter-spacing:.08em;text-transform:uppercase;margin-bottom:7px}}
  .doha[hidden]{{display:none}}
  .site-disclaimer{{padding:25px max(24px,calc((100% - 1180px)/2));border-top:1px solid var(--line);color:var(--muted);font-size:11px}}
  .site-disclaimer strong{{display:block;color:var(--ink);font:500 10px var(--mono);letter-spacing:.08em;text-transform:uppercase;margin-bottom:7px}}
  @media(max-width:860px){{.doha-list{{grid-template-columns:1fr}}}}
  @media(max-width:700px){{
    .doha-hero,.doha-tools,.doha-disclaimer,.site-disclaimer{{padding-left:16px;padding-right:16px}}
    .doha-list,.doha-note{{padding-left:16px;padding-right:16px}}
    .open-all{{margin-left:0}}
    .couplet{{font-size:18px}}
  }}
</style>
</head>
<body>
{strip}
  <main>
    <section class="doha-hero"><p class="eyebrow">कबीर • रहीम • तुलसीदास</p><h1>निवेश और जीवन के 100 दोहे</h1><p>{e(intro)}</p></section>
    <section class="doha-tools">
      <input type="search" id="doha-search" placeholder="खोजें — जैसे SIP, धैर्य, Emergency, 42" aria-label="दोहे खोजें">
      <div class="poet-chips"><button type="button" class="active" data-poet="">सभी <small>{len(dohas)}</small></button>{chips}</div>
      <span class="doha-count" id="doha-count">{len(dohas)} दोहे</span>
      <button type="button" class="open-all" id="open-all">सब खोलें</button>
    </section>
    <section class="doha-list" id="doha-list">
{chr(10).join(cards)}
    </section>
    <section class="doha-note"><h2>निष्कर्ष</h2><p>{e(outro)}</p></section>
  </main>
  <aside class="doha-disclaimer"><strong>Disclaimer</strong>यह पेज केवल शैक्षिक उद्देश्य के लिए है, यह निवेश सलाह नहीं है। उदाहरणों के आंकड़े अनुमानित हैं; पिछला प्रदर्शन भविष्य की गारंटी नहीं है। किसी कंपनी या fund का नाम केवल उदाहरण के लिए है, खरीदने-बेचने की सलाह नहीं। Mutual Fund investments are subject to market risks, read all scheme related documents carefully. Ratnesh Kumar Singh — NISM certified Mutual Fund Distributor (AMFI ARN-132137). Not a SEBI-registered Research Analyst. निवेश से पहले अपने वित्तीय सलाहकार से परामर्श करें।</aside>
{tail}
  <script src="nav.js"></script>
  <script>
  (() => {{
    const cards = [...document.querySelectorAll('.doha')];
    const search = document.getElementById('doha-search');
    const count = document.getElementById('doha-count');
    const chips = [...document.querySelectorAll('.poet-chips button')];
    const openAll = document.getElementById('open-all');
    let poet = '';
    function filter() {{
      const q = search.value.trim().toLowerCase();
      let shown = 0;
      for (const c of cards) {{
        const ok = (!poet || c.dataset.poet === poet) &&
          (!q || c.id === 'doha-' + q || c.textContent.toLowerCase().includes(q));
        c.hidden = !ok; if (ok) shown++;
      }}
      count.textContent = shown + ' दोहे';
    }}
    search.addEventListener('input', filter);
    chips.forEach(b => b.addEventListener('click', () => {{
      chips.forEach(x => x.classList.toggle('active', x === b));
      poet = b.dataset.poet; filter();
    }}));
    openAll.addEventListener('click', () => {{
      const open = openAll.textContent === 'सब खोलें';
      document.querySelectorAll('.doha details').forEach(d => d.open = open);
      openAll.textContent = open ? 'सब बंद करें' : 'सब खोलें';
    }});
  }})();
  </script>
</body>
</html>
"""
OUT.write_text(page, encoding="utf-8")
print("wrote", OUT, len(dohas), "dohas")
